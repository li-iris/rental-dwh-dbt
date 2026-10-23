"""
Step 1 - Bronze extraction.

Reads every .xlsx in a folder and writes ONE csv per sheet, plus an inventory csv.
Rules (this is the "Bronze" idea):
  * Nothing is cleaned, renamed or interpreted. Every value is stored as text.
  * Each row keeps where it came from: file, file hash, sheet, original row number.
  * Columns are named by their Excel letter (A, B, ... AA) because the real
    headers are messy / multi-row. Header cleaning happens later (Silver).
  * Same input -> same output (re-running never creates duplicates).

Usage:  python scripts/01_extract_to_csv.py raw_private
Output: raw_private/bronze_csv/*.csv  and  raw_private/bronze_csv/_inventory.csv
        (inside raw_private, so it is git-ignored and never uploaded)
"""
import csv
import hashlib
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def safe(name: str) -> str:
    """Make a string safe for a file name (keeps Chinese characters)."""
    cleaned = re.sub(r"[^0-9A-Za-z一-鿿]+", "_", name).strip("_")
    return cleaned[:60] or "x"


def to_text(value) -> str:
    """Convert any cell value to text without changing its meaning."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def count_formulas(ws_formula, ws_value):
    """Return (formula cells, formula cells whose calculated value is missing)."""
    formulas = 0
    no_cache = 0
    for row in ws_formula.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                formulas += 1
                if is_blank(ws_value[cell.coordinate].value):
                    no_cache += 1
    return formulas, no_cache


def main(raw_dir: str) -> None:
    raw = Path(raw_dir)
    out = raw / "bronze_csv"
    out.mkdir(parents=True, exist_ok=True)

    files = sorted(p for p in raw.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        sys.exit(f"No .xlsx files found in {raw}")

    inventory = []
    for path in files:
        file_hash = sha256_of(path)
        wb_val = load_workbook(path, data_only=True)    # calculated values
        wb_fml = load_workbook(path, data_only=False)   # formulas
        stem = safe(path.stem)

        for idx, ws in enumerate(wb_val.worksheets, start=1):
            rows = []          # (original_row_number, [text values])
            width = 0
            nonempty_cells = 0
            for r_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if all(is_blank(v) for v in row):
                    continue
                last = max(i for i, v in enumerate(row, start=1) if not is_blank(v))
                width = max(width, last)
                nonempty_cells += sum(1 for v in row if not is_blank(v))
                rows.append((r_no, [to_text(v) for v in row]))

            out_name = f"{stem}__{idx:02d}_{safe(ws.title)}.csv"
            with open(out / out_name, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(["_source_file", "_file_sha256", "_sheet", "_row_number"]
                           + [get_column_letter(i) for i in range(1, width + 1)])
                for r_no, vals in rows:
                    vals = (vals + [""] * width)[:width]
                    w.writerow([path.name, file_hash, ws.title, r_no] + vals)

            fml, no_cache = count_formulas(wb_fml[ws.title], ws)
            inventory.append({
                "file": path.name,
                "sheet": ws.title,
                "state": ws.sheet_state,
                "reported_max_row": ws.max_row,
                "nonempty_rows": len(rows),
                "width_cols": width,
                "nonempty_cells": nonempty_cells,
                "formula_cells": fml,
                "formula_no_cached_value": no_cache,
                "csv": out_name,
            })

    with open(out / "_inventory.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(inventory[0].keys()))
        w.writeheader()
        w.writerows(inventory)

    print(f"{len(files)} workbooks, {len(inventory)} sheets -> {out}")
    for item in inventory:
        print(f"{item['nonempty_rows']:5d} rows x {item['width_cols']:3d} cols | "
              f"formulas {item['formula_cells']:4d} (no value: {item['formula_no_cached_value']}) | "
              f"{item['sheet']}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "raw_private")
