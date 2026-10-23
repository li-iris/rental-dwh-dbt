"""Generate FULLY SYNTHETIC Excel workbooks that have the same sheet titles and column positions as the real ones.

Nothing here comes from a real business: streets, people and amounts are invented. The unchanged pipeline
(01_extract_to_csv -> 02_load_bronze -> dbt) runs on these files, so anyone can reproduce the project.

Some rows are broken ON PURPOSE (overdue water bills, unclear rents, expired parking, too many inspections ...)
so the data-quality tests and the problem list have something to find.

Usage:  python scripts/gen_synthetic.py [output_dir]      (default: raw_private)
"""
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

random.seed(42)
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "raw_private")
OUT.mkdir(parents=True, exist_ok=True)

STREETS = ["Marlow St", "Kestrel Rd", "Harbourview Ave", "Linden St", "Orchard Pde"]
STAFF = ["Sam", "Jordan", "Casey"]
FIRST = ["Alex", "Mia", "Noah", "Lena", "Omar", "Priya", "Kai", "Sofia", "Ethan", "Yuki", "Liam", "Aria", "Ravi", "Zoe", "Hugo"]
LAST = ["Anders", "Brook", "Castle", "Dunn", "Ellis", "Frost", "Grant", "Hale", "Ivory", "Joyce"]
VISA = ["Student 500", "Citizen", "PR", "Working holiday 417", "Graduate 485"]
ROOMS = ["Room A", "Room B", "Room C", "Master", "Living room"]


def d(y, m, dd):
    return date(y, m, dd)


def dot(x):          # 2025.6.10
    return f"{x.year}.{x.month}.{x.day}"


def md(x):           # 6.18
    return f"{x.month}.{x.day}"


def person():
    return f"{random.choice(FIRST)} {random.choice(LAST)}"


# ---------- invented properties ----------
PROPS = []
for i in range(1, 19):
    street = STREETS[i % len(STREETS)]
    unit = random.randint(1, 40) if i % 3 else None
    number = random.randint(2, 180)
    addr = f"{unit}/{number} {street}" if unit else f"{number} {street}"
    PROPS.append({"addr": addr, "rooms": random.randint(2, 4)})


def put(ws, row, col, value):
    if value not in (None, ""):
        ws.cell(row=row, column=col, value=value)


def letters_to_n(s):
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def row_of(ws, r, cells):
    """cells: {'A': value, 'AC': value}"""
    for k, v in cells.items():
        put(ws, r, letters_to_n(k), v)


def pad_header(ws, width, row=1):
    """Make sure the sheet is `width` columns wide (a header row with a placeholder in the last column)."""
    if ws.cell(row=row, column=width).value is None:
        ws.cell(row=row, column=width, value=".")


# ---------- sheets ----------
def sheet_tenant(wb, props, snapshot):
    ws = wb.create_sheet("0.Tenant")
    row_of(ws, 1, {"A": "Address", "B": "Room", "C": "Tenant", "D": "Lease term", "E": "Rent (pw)", "K": "Visa"})
    r = 2
    for pi, p in enumerate(props):
        for k in range(p["rooms"]):
            start = d(2025, random.randint(1, 8), random.randint(1, 28))
            end = start.replace(year=start.year + 1) - timedelta(days=1)
            rent = random.choice([320, 350, 380, 420, 450, 480, 520])
            rent_cell = str(rent)
            # deliberate messy cells
            if (pi + k) % 11 == 0:
                rent_cell = f"{rent}, from {md(start + timedelta(days=120))} {rent + 20}"
            elif (pi + k) % 13 == 0:
                rent_cell = f"{rent}-{rent + 30}"
            elif (pi + k) % 17 == 0:
                rent_cell = "negotiable"
            term = f"{dot(start)}-{dot(end)}"
            if (pi + k) % 19 == 0:
                term = "ongoing"
            if snapshot == "old" and pi >= 14:
                continue
            row_of(ws, r, {"A": p["addr"], "B": ROOMS[k], "C": person(), "D": term, "E": rent_cell,
                           "K": random.choice(VISA)})
            r += 1
    pad_header(ws, 11)


def sheet_water(wb, props):
    ws = wb.create_sheet("5.3水费结算")
    row_of(ws, 1, {"A": "Address", "B": "Note", "C": "Period", "D": "Amount", "E": "Bill copy", "F": "Company paid",
                   "G": "Invoice sent", "H": "Tenant paid", "I": "Checked"})
    r = 2
    for pi, p in enumerate(props[:12]):
        first = True
        for q in range(3):
            ps = d(2025, 1 + q * 3, 1)
            pe = (d(2025, 4 + q * 3, 1) - timedelta(days=1))
            paid = pe + timedelta(days=random.randint(5, 25))
            cells = {"C": f"{dot(ps)}-{dot(pe)}", "D": random.choice([88.4, 120.0, 143.25, 96.8, 210.5]),
                     "E": "有", "F": f"{random.choice(STAFF)}{dot(paid)}"}
            if first:
                cells["A"] = p["addr"]
                first = False
            # deliberate states
            kind = (pi + q) % 5
            if kind in (0, 1):      # fully settled
                inv = paid + timedelta(days=random.randint(1, 10))
                cells["G"] = f"{md(inv)}{random.choice(STAFF).lower()}"
                cells["H"] = f"{md(inv + timedelta(days=7))}{random.choice(STAFF).lower()}"
            elif kind == 2:         # invoice sent, tenant has not paid
                inv = paid + timedelta(days=4)
                cells["G"] = f"{md(inv)}{random.choice(STAFF).lower()}"
            elif kind == 3:         # company paid, no invoice yet  -> overdue / may be unrecoverable
                pass
            else:                   # settled in the exit report
                cells["G"] = "退房报告"
            row_of(ws, r, cells)
            r += 1
    pad_header(ws, 9)


def sheet_utility(wb, props):
    ws = wb.create_sheet("5.2电和煤气bill结算")
    row_of(ws, 1, {"A": "Address", "B": "Type", "C": "Period", "D": "Due", "E": "Amount", "F": "Calc file",
                   "G": "Invoice sent", "H": "Payment check"})
    r = 2
    for pi, p in enumerate(props[:10]):
        ps, pe = d(2025, 7, 1), d(2025, 9, 30)
        cells = {"A": p["addr"], "B": random.choice(["电", "煤气", "电+煤气"]),
                 "C": f"{dot(ps)}-{dot(pe)}", "D": dot(pe + timedelta(days=14)),
                 "E": random.choice(["180", "240.5", "315", "96.2"]), "F": "calc.xlsx"}
        if pi % 4 == 0:
            cells["E"] = "600(manager said charge 600)"
        if pi % 3 != 2:
            cells["G"] = f"{md(pe + timedelta(days=10))}{random.choice(STAFF).lower()}"
        if pi % 3 == 0:
            cells["H"] = f"{md(pe + timedelta(days=20))}{random.choice(STAFF).lower()}"
        if pi == 9:
            cells = {"A": p["addr"], "B": "电"}      # bill type only: not set up yet
        row_of(ws, r, cells)
        r += 1
    pad_header(ws, 8)


def sheet_inspection(wb, props):
    ws = wb.create_sheet("5.1查房和消防通知")
    row_of(ws, 1, {"A": "Routine inspection", "M": "Fire safety", "U": "Other notices"})
    row_of(ws, 2, {"A": "Address", "B": "Date", "D": "Time", "E": "Email", "F": "Message", "G": "Pre-check",
                   "H": "Pre-check by", "I": "Attended", "M": "Address", "N": "Date", "P": "Time", "Q": "Email",
                   "R": "Done", "U": "Address", "V": "Date", "X": "Time", "Y": "Category", "Z": "Notice date",
                   "AA": "Email", "AB": "Follow up", "AC": "Done"})
    r = 3
    for pi, p in enumerate(props[:14]):
        n = 5 if pi == 3 else random.randint(1, 3)           # property 3: too many inspections in a year
        for k in range(n):
            day = d(2025, 1 + k * 2, random.randint(1, 20))
            row_of(ws, r, {"A": p["addr"], "B": dot(day), "D": "10:00", "E": f"已发 {random.choice(STAFF)}",
                           "F": "已发", "G": "ok", "H": random.choice(STAFF), "I": random.choice(STAFF)})
            if pi % 4 == 0:
                row_of(ws, r, {"M": p["addr"], "N": dot(day + timedelta(days=3)), "P": "14:00",
                               "Q": f"已发 {random.choice(STAFF)}", "R": "完成" if k == 0 else None})
            if pi % 5 == 0 and k == 0:
                row_of(ws, r, {"U": p["addr"], "V": dot(day + timedelta(days=5)), "X": "09:00", "Y": "Smoke alarm",
                               "Z": dot(day - timedelta(days=2)), "AA": f"已发 {random.choice(STAFF)}",
                               "AB": random.choice(STAFF), "AC": "done"})
            r += 1
    pad_header(ws, 29, row=1)


def sheet_special(wb, props):
    ws = wb.create_sheet("2.1房子特别条款&情况登记")
    row_of(ws, 1, {"A": "Address", "B": "Special term", "C": "Carried out", "D": "Recorded"})
    terms = [("Tenant may keep one small dog", "approved"), ("Landlord repaints lounge before move-in", "done"),
             ("Rent reduced by $20 pw for 3 months", "$240 total, applied"), ("No smoking on balcony", None),
             ("Replace oven within 30 days", None)]
    r = 2
    for pi, p in enumerate(props[:10]):
        t, s = terms[pi % len(terms)]
        row_of(ws, r, {"A": p["addr"], "B": t, "C": s, "D": f"{md(d(2025, 5, 1) + timedelta(days=pi * 9))} {random.choice(STAFF)}"})
        r += 1
    pad_header(ws, 21)


def sheet_garage(wb, props):
    ws = wb.create_sheet("车库车位信息表")
    row_of(ws, 1, {"A": "Parking list", "E": "Garage list", "J": "No tenant", "P": "Monthly", "Q": "Price",
                   "R": "Bond", "S": "Access", "T": "Cover", "X": "Weekly", "Y": "Price", "Z": "Bond", "AA": "Access"})
    rows = [
        ({"A": "Marlow St 12/3 P1", "B": "合同已到期"}, {}),
        ({"A": "Kestrel Rd 45 P2", "B": dot(d(2026, 3, 31))}, {}),
        ({"A": "Linden St 9/2 P3", "B": dot(d(2025, 9, 30))}, {"E": "Linden St 9/2 Garage 1", "F": "合同已到期"}),
        ({}, {"E": "Orchard Pde 70 Garage", "F": dot(d(2026, 8, 31))}),
        ({}, {"J": "Harbourview Ave 31 P4", "K": "没有租客消息"}),
    ]
    prices = [("Marlow St 12", 260, 520, "Remote", "Covered"), ("Kestrel Rd 45", 240, 480, "Key", "Open"),
              ("Linden St 9", 300, 600, "Remote", "Covered")]
    r = 2
    for i, (a, b) in enumerate(rows):
        cells = dict(a)
        cells.update(b)
        if i < len(prices):
            n, pr, bd, ac, cv = prices[i]
            cells.update({"P": n, "Q": pr, "R": bd, "S": ac, "T": cv})
        if i < 2:
            cells.update({"X": f"{STREETS[i]} {10 + i}".replace("Marlow St", "Marlow St"), "Y": 70 + i * 10, "Z": 140, "AA": "Key"})
        row_of(ws, r, cells)
        r += 1
    pad_header(ws, 27)


def sheet_moving(wb, props):
    ws = wb.create_sheet("搬家工作记录")
    items = ["Bed", "Mattress", "Desk", "Chair", "Wardrobe", "Sofa", "Microvave", "Boxes", "Fridge", "TV", "Bike", "Lamp", "Shelf", "Others"]
    row_of(ws, 1, {"I": items[0]})
    for j, name in enumerate(items):
        put(ws, 1, 9 + j, name)
    row_of(ws, 2, {"A": "Date", "B": "Address", "C": "Rooms", "D": "Stairs", "E": "Hours", "F": "Others", "G": "Crew",
                   "W": "Start", "X": "End", "Y": "Note", "Z": "Problem"})
    r = 3
    for i in range(10):
        p = PROPS[i]
        day = d(2025, random.randint(2, 9), random.randint(1, 27))
        date_cell = day.isoformat() if i % 3 == 0 else (f"{day.day}/{day.month}/{day.year}" if i % 3 == 1 else md(day))
        cells = {"A": date_cell, "B": p["addr"], "C": f"{random.randint(1, 3)}b", "D": random.choice(["no", "1 flight", "2 flights"]),
                 "E": random.choice(["4", "8.5+4.5", "6+3", "about 5"]), "G": " ".join(random.sample(STAFF, 2)),
                 "W": "08:00", "X": "15:00"}
        for j in range(14):
            if random.random() < 0.5:
                cells[get_column_letter(9 + j)] = random.randint(0, 6)
        if i == 4:
            cells["Z"] = "Lift was out of service"
        row_of(ws, r, cells)
        r += 1
    pad_header(ws, 26)


def sheet_rent_cash(wb, props):
    ws = wb.create_sheet("房租cash")
    row_of(ws, 1, {"AO": "10月", "BA": "11月"})
    pad_header(ws, 61)
    r = 2
    for p in props[:6]:
        for k in range(min(p["rooms"], 3)):
            c = {"U": ROOMS[k]}
            if k == 0:
                c["D"] = p["addr"]
            col = "AO" if k % 2 == 0 else "BA"
            base = letters_to_n(col)
            col_at = lambda off: get_column_letter(base + off)
            t0 = d(2025, 10 if col == "AO" else 11, random.randint(1, 8))
            c[col_at(0)] = random.choice([320, 380, 450])
            c[col_at(1)] = random.choice(STAFF)
            c[col_at(2)] = t0.isoformat()
            if r % 3:
                c[col_at(4)] = (t0 + timedelta(days=4)).isoformat()
                c[col_at(3)] = 4
                if r % 2:
                    c[col_at(6)] = (t0 + timedelta(days=6)).isoformat()
                    c[col_at(5)] = 2
                    c[col_at(7)] = random.choice(STAFF)
            row_of(ws, r, c)
            r += 1


def empty_sheet(wb, title, width):
    ws = wb.create_sheet(title)
    for i in range(1, width + 1):
        ws.cell(row=1, column=i, value="col " + get_column_letter(i))


SIMPLE_EMPTY = {"Profit": 64, "rent credit": 64, "人民币入账": 64, "所有房子入账核对": 64, "1.招租+入住+退租list": 64,
                "2.办理租房手续list": 64, "2. 财务日常入账": 64, "3.转合同-变更信息表格": 64,
                "4.办理续租Rental review list": 64, "房子物品list": 64}


def workbook(path, snapshot):
    wb = Workbook()
    wb.remove(wb.active)
    props = PROPS if snapshot == "new" else PROPS[:16]
    sheet_tenant(wb, props, snapshot)
    sheet_water(wb, props)
    sheet_utility(wb, props)
    sheet_inspection(wb, props)
    sheet_special(wb, props)
    sheet_garage(wb, props)
    sheet_moving(wb, props)
    sheet_rent_cash(wb, props)
    for t, w in SIMPLE_EMPTY.items():
        empty_sheet(wb, t, w)
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    workbook(OUT / "Sample_working_list_2025-06.xlsx", "old")
    workbook(OUT / "Sample_working_list.xlsx", "new")
