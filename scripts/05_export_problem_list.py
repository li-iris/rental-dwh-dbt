"""Export the problem list: every follow-up the automatic checks found, one sheet per kind of problem.

dbt stores the failing rows of every check when it runs `dbt test --store-failures`. This script reads those stored
rows, adds a plain-language description (English and Chinese) and two empty columns, "Owner answer" and "Done",
and writes reports/problem_list.xlsx. Fill in the answers in Excel, fix the source data, refresh, and solved
problems disappear on their own.

Run:  python scripts/05_export_problem_list.py            (after dbt test --store-failures)
"""
import os
import datetime
import decimal
import sys
from pathlib import Path

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DB = "dbname=" + (os.environ.get("RENTAL_DB") or os.environ.get("PGDATABASE") or "rental_dwh")
OUT = Path(__file__).resolve().parent.parent / "reports" / "problem_list.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ANSWER_FILL = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")

# test name -> (area, sheet name, what is wrong (EN), 中文说明, what to do (EN))
# Tests that are not listed here still appear, with their raw name, so nothing is hidden.
CATALOG = {
    "warn_water_bill_company_paid_invoice_not_sent": ("Water", "Water-invoice overdue", "The agency paid the water bill but no invoice to the tenant is recorded within the allowed days.", "公司已垫付水费,但超过规定天数仍没有给租客发发票的记录。", "Send the invoice and record the date, or record that it was settled in the exit report."),
    "warn_water_bill_may_be_unrecoverable": ("Water", "Water-may be lost", "Invoice still not sent after the 3-month NSW window, so the tenant may no longer have to pay.", "超过新州规定的 3 个月期限还没发发票,租客可能不再需要支付。", "Decide: bill it anyway, ask the landlord, or write it off."),
    "warn_water_bill_tenant_payment_not_recorded": ("Water", "Water-payment overdue", "Invoice was sent but no tenant payment is recorded within the allowed days.", "发票已发出,但超过规定天数仍没有租客付款记录。", "Chase the tenant and record the payment date."),
    "warn_water_bill_period_unreadable": ("Water", "Water-period unreadable", "The billing period cell could not be read.", "账单期间读不出来。", "Correct the period in the source."),
    "warn_utility_bill_not_set_up": ("Utilities", "Utility-not set up", "Electricity/gas is on the list but nothing is filled in.", "电费/煤气在清单上,但什么都没填。", "Set up the account or remove the line."),
    "warn_utility_bill_invoice_not_sent": ("Utilities", "Utility-invoice not sent", "A bill has an amount but no invoice is recorded.", "有账单金额,但没有发票记录。", "Send the invoice and record it."),
    "warn_rent_roll_rent_unclear": ("Rent", "Rent-unclear", "The weekly rent for this property-month cannot be worked out from the leases.", "根据租约算不出这个房产这个月的应收租金。", "Check the leases of this property; fill in the missing rent."),
    "warn_rent_roll_recorded_differs_from_expected": ("Rent", "Rent-differs", "Recorded rent received differs from the expected rent by more than the allowed amount.", "记录的收款与应收租金相差超过允许范围。", "Check the receipt records or the lease rent."),
    "warn_rent_roll_credit_exceeds_expected": ("Rent", "Rent-credit too big", "A vacancy credit is larger than the expected rent.", "空置抵扣比应收租金还大。", "Check the credit amount and its dates."),
    "warn_rent_payment_property_has_no_tenancy": ("Rent", "Payment-no lease", "A rent payment was entered for a property with no lease covering that date.", "录入了一笔租金,但该房产在这个日期没有对应租约。", "Enter the lease or correct the address/date."),
    "warn_account_check_issue_to_follow_up": ("Rent", "Account check issues", "A rent account check ended as an issue.", "租金账户核对结果是“有问题”。", "Resolve the issue noted in the check."),
    "warn_bank_ledger_days_written_mismatch": ("Rent", "Ledger days mismatch", "Days written in the bank ledger differ from the days between dates.", "银行台账里写的天数与日期算出的天数不一致。", "Correct the days or the dates."),
    "warn_bank_ledger_gap_between_periods": ("Rent", "Ledger gaps", "A gap between two bank ledger periods.", "银行台账两个时间段之间有空档。", "Find the missing period."),
    "warn_process_step_done_without_date": ("Paperwork", "Step done-no date", "A paperwork step is marked done but has no date.", "手续步骤标了已完成,但没有日期。", "Add the date."),
    "warn_process_step_text_needs_review": ("Paperwork", "Step text unclear", "A step cell holds text that cannot be read as done or not done.", "步骤格子里只有文字,看不出是否完成。", "Replace with a clear yes/date."),
    "warn_process_case_without_anchor_date": ("Paperwork", "Case-no anchor date", "A paperwork case has no contract end or move-in date, so deadlines cannot be calculated.", "手续案例没有合同到期日或入住日,算不出截止时间。", "Add the date."),
    "warn_process_case_address_not_matched": ("Paperwork", "Case-address unknown", "The address matches no managed property.", "地址对不上任何在管房产。", "Fix the address or add the property."),
    "warn_inspection_address_not_matched": ("Inspections", "Inspection-address unknown", "The address matches no managed property.", "地址对不上任何在管房产。", "Fix the address or add the property."),
    "warn_inspection_date_unreadable": ("Inspections", "Inspection-date unreadable", "The inspection date could not be read.", "查房日期读不出来。", "Correct the date."),
    "warn_routine_inspections_exceed_yearly_limit": ("Inspections", "Too many inspections", "More routine inspections in 12 months than the allowed number.", "12 个月内例行检查次数超过允许的次数。", "Move or cancel the extra inspection."),
    "warn_routine_inspection_notice_too_short": ("Inspections", "Inspection notice short", "The tenant was told too few days before a routine inspection.", "例行检查前通知租客的天数不够。", "Re-issue the notice with enough days."),
    "warn_moving_job_date_incomplete": ("Moving", "Moving-date incomplete", "The moving date has no year or is empty.", "搬家日期缺年份或为空。", "Write the full date."),
    "warn_moving_job_without_item_counts": ("Moving", "Moving-no items", "A moving job has no item counts.", "搬家任务没有物品数量。", "Enter the items moved."),
    "warn_bill_or_moving_address_not_matched": ("Moving", "Address unknown (bills)", "A bill or moving job address matches no managed property.", "账单或搬家的地址对不上任何在管房产。", "Fix the address or add the property."),
    "warn_special_term_without_status": ("Special terms", "Clause-no status", "A special clause has no record of being carried out.", "特别条款没有落实情况的记录。", "Record what happened."),
    "warn_special_term_address_not_matched": ("Special terms", "Clause-address unknown", "The address matches no managed property.", "地址对不上任何在管房产。", "Fix the address or add the property."),
    "warn_parking_space_not_linked_to_property": ("Parking", "Parking-no property", "A parking space cannot be linked to a managed property.", "车位对不上任何在管房产。", "Fix the address or leave it if it is not ours."),
    "warn_turnover_case_without_unique_tenancy": ("Move-out", "Case-lease unclear", "A move-out case cannot be matched to exactly one lease.", "搬出案例对不上唯一一份租约。", "Check which lease it belongs to."),
    "warn_person_name_not_in_staff_list": ("Data quality", "Unknown staff name", "A staff name is not in the staff list.", "员工名字不在员工名单里。", "Add the spelling to staff_aliases."),
}
AREA_ORDER = ["Water", "Rent", "Utilities", "Inspections", "Paperwork", "Move-out", "Moving", "Special terms", "Parking", "Data quality", "Other"]


def clean(v):
    if isinstance(v, decimal.Decimal):
        return float(v)
    if isinstance(v, datetime.datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    if isinstance(v, (dict, list)):
        return str(v)
    return v


def main():
    with psycopg.connect(DB) as conn, conn.cursor() as cur:
        cur.execute("""select table_schema, table_name from information_schema.tables
                       where table_schema like '%dbt_test__audit%' and table_name ~ '^(warn|assert)_' order by 2""")
        found = cur.fetchall()
        if not found:
            sys.exit("No stored test failures found. Run:  dbt test --store-failures   (run_all.sh does this).")
        problems = []
        for schema, name in found:
            cur.execute(f'select * from "{schema}"."{name}"')
            cols = [d.name for d in cur.description]
            rows = [[clean(v) for v in r] for r in cur.fetchall()]
            if rows:
                problems.append((name, cols, rows))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    errors = [p for p in problems if p[0].startswith("assert_")]
    ws.append(["Problem list - generated " + datetime.date.today().isoformat()])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Each row below is one kind of problem; open its sheet to see every record. Write your answer in the yellow "
               "'Owner answer' column of that sheet. Fix the source data (Excel or Navicat entry), refresh, and solved items disappear."])
    ws.append([])
    ws.append(["Area", "Sheet", "How many", "What is wrong", "中文说明", "What to do", "Check name"])
    for j in range(1, 8):
        c = ws.cell(row=4, column=j)
        c.fill, c.font = HEADER_FILL, HEADER_FONT

    def sort_key(p):
        area = CATALOG.get(p[0], ("Other",))[0]
        return (AREA_ORDER.index(area) if area in AREA_ORDER else len(AREA_ORDER), -len(p[2]))

    used = set()
    detail = []
    for name, cols, rows in sorted(problems, key=sort_key):
        area, sheet, en, zh, todo = CATALOG.get(name, ("Other", name[:28], name, "", ""))
        base = sheet[:28]
        sheet_name, k = base, 2
        while sheet_name in used:
            sheet_name, k = f"{base[:26]}{k}", k + 1
        used.add(sheet_name)
        is_error = name.startswith("assert_")
        ws.append([area, sheet_name, len(rows), en, zh, todo, name])
        r = ws.max_row
        ws.cell(row=r, column=2).hyperlink = f"#'{sheet_name}'!A1"
        ws.cell(row=r, column=2).font = Font(color="0563C1", underline="single")
        if is_error:
            for j in range(1, 8):
                ws.cell(row=r, column=j).fill = RED
        detail.append((sheet_name, cols, rows, name, en, zh))
    if not problems:
        ws.append(["", "No problems found.", 0])
    for j, w in enumerate([14, 26, 10, 60, 40, 50, 46], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"

    for sheet_name, cols, rows, name, en, zh in detail:
        s = wb.create_sheet(sheet_name)
        s.append([f"{en}  {zh}  ({name})"])
        s["A1"].font = Font(bold=True)
        headers = cols + ["Owner answer", "Done"]
        s.append(headers)
        for j in range(1, len(headers) + 1):
            c = s.cell(row=2, column=j)
            c.fill, c.font = (ANSWER_FILL, Font(bold=True)) if j > len(cols) else (HEADER_FILL, HEADER_FONT)
        for row in rows:
            s.append(row + [None, None])
        for j, h in enumerate(headers, start=1):
            longest = max([len(str(h))] + [len(str(r[j - 1])) for r in rows[:200] if j - 1 < len(r) and r[j - 1] is not None])
            s.column_dimensions[get_column_letter(j)].width = min(max(12, longest + 2), 60)
        for row in s.iter_rows(min_row=3):
            for c in row:
                if isinstance(c.value, (datetime.date, datetime.datetime)):
                    c.number_format = "yyyy-mm-dd"
                c.alignment = Alignment(vertical="top", wrap_text=True)
        s.freeze_panes = "B3"
        s.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(rows) + 2}"

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}  ({len(problems)} kinds of problem, {sum(len(p[2]) for p in problems)} records, {len(errors)} of them errors)")


if __name__ == "__main__":
    main()
