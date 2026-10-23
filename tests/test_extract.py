"""Step 1: Excel -> bronze CSV. Nothing may be cleaned or lost; re-running must give the same result."""
import csv
from datetime import date, datetime

from openpyxl import Workbook

from helpers import load_script

extract = load_script("01_extract_to_csv.py")


def test_to_text_keeps_meaning():
    assert extract.to_text(None) == ""
    assert extract.to_text(3.0) == "3"          # 3.0 from Excel is the number 3
    assert extract.to_text(3.5) == "3.5"
    assert extract.to_text(True) == "TRUE"       # bool must not become "1"
    assert extract.to_text(date(2025, 6, 1)) == "2025-06-01"
    assert extract.to_text(datetime(2025, 6, 1, 9, 30)) == "2025-06-01 09:30:00"


def test_is_blank():
    assert extract.is_blank(None)
    assert extract.is_blank("   ")
    assert not extract.is_blank(0)               # zero is a value, not blank


def test_safe_file_name_keeps_chinese_and_strips_symbols():
    assert extract.safe("5.3水费结算") == "5_3水费结算"
    assert extract.safe("???") == "x"
    assert len(extract.safe("a" * 200)) == 60


def test_sha256_changes_with_content(tmp_path):
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_bytes(b"hello")
    b.write_bytes(b"hello!")
    assert extract.sha256_of(a) == extract.sha256_of(a)
    assert extract.sha256_of(a) != extract.sha256_of(b)


def _make_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "0.Tenant"
    ws["A1"] = "Name"
    ws["B1"] = "Rent"
    ws["A3"] = "Tenant A"            # row 2 is empty on purpose: empty rows are skipped, row numbers are kept
    ws["B3"] = 450.0
    ws["C3"] = "=B3*2"
    wb.create_sheet("Empty")
    wb.save(path)


def test_main_writes_one_csv_per_sheet_with_lineage_columns(tmp_path):
    _make_workbook(tmp_path / "book.xlsx")
    extract.main(str(tmp_path))
    out = tmp_path / "bronze_csv"
    files = sorted(p.name for p in out.glob("*.csv"))
    assert "_inventory.csv" in files
    assert len(files) == 3                        # 2 sheets + inventory

    tenant = next(out.glob("book__01_*.csv"))
    rows = list(csv.reader(open(tenant, encoding="utf-8-sig")))
    assert rows[0][:4] == ["_source_file", "_file_sha256", "_sheet", "_row_number"]
    assert rows[0][4:] == ["A", "B"]          # formula C3 has no cached value (file never opened in Excel), so it counts as blank
    assert [r[3] for r in rows[1:]] == ["1", "3"]  # original row numbers kept, blank row 2 skipped
    assert rows[2][4:6] == ["Tenant A", "450"]
    assert rows[1][0] == "book.xlsx"


def test_main_is_repeatable(tmp_path):
    _make_workbook(tmp_path / "book.xlsx")
    extract.main(str(tmp_path))
    first = {p.name: p.read_bytes() for p in (tmp_path / "bronze_csv").glob("*.csv")}
    extract.main(str(tmp_path))
    second = {p.name: p.read_bytes() for p in (tmp_path / "bronze_csv").glob("*.csv")}
    assert first == second


def test_inventory_flags_formulas_without_cached_value(tmp_path):
    _make_workbook(tmp_path / "book.xlsx")
    extract.main(str(tmp_path))
    inv = list(csv.DictReader(open(tmp_path / "bronze_csv" / "_inventory.csv", encoding="utf-8-sig")))
    tenant = next(r for r in inv if r["sheet"] == "0.Tenant")
    assert tenant["formula_cells"] == "1"
    assert tenant["formula_no_cached_value"] == "1"
